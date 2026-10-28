"""Karnataka ADB EIA as a third report: its own site profile (nothing from Great Nicobar or
Cuddalore), the statewide species-list rule, checks that are "not applicable to this report",
its own pairing places, the status-table headers it uses, the Kannada letter and the lender
addressee. No model, GBIF or Tavily call is made (all mocked or refused)."""

import io
import json
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import export_results as ex  # noqa: E402
from tigerfish import llm, pairing, sites, species, status_pages  # noqa: E402
from tigerfish.ui import letter, rti  # noqa: E402
from tests.test_letter import FINDINGS, doc_text  # noqa: E402
from tests.test_species import BIRD, MATCH  # noqa: E402

KA = sites.profile(sites.KARNATAKA)
KA_REPORT = {"title": "Karnataka Rural Roads Improvement Project EIA (ADB, July 2026)",
             "source": "EIA prepared by KRRDA for the Asian Development Bank", "state": "Karnataka",
             "letter_to": KA["letter_to"],
             "not_applicable": {k: f"{sites.NOT_APPLICABLE}: {v}"
                                for k, v in KA["not_applicable"].items()}}


# ---------- site profile ----------

def test_profile_is_karnataka_and_nothing_carries_over():
    assert KA["state"] == "Karnataka"
    assert KA["archipelago_gadm"] is None                  # the A&N archipelago rule is off
    assert KA["region_gadm"] == "IND.16_1"                 # GBIF GADM "Karnataka"
    assert KA["consultant"] is None                        # no consultant firm named
    lat, lon = KA["site"]
    assert 11.5 <= lat <= 18.5 and 74 <= lon <= 78.5       # inside the extent on report p.23
    for other in (sites.GREAT_NICOBAR, sites.CUDDALORE):
        o = sites.profile(other)
        assert o["site"] != KA["site"] and o["state"] != KA["state"]
        assert not o.get("not_applicable") and not o.get("pairing") and not o.get("region_gadm")
    text = json.dumps(KA)
    for word in ("Nicobar", "Galathea", "Cuddalore", "Tamil", "Vimta", "Falcon"):
        assert word not in text


def test_not_applicable_checks():
    for check in ("omission", "accreditation", "promise", "public_consultation"):
        assert sites.not_applicable(sites.KARNATAKA, check)
    assert sites.not_applicable(sites.KARNATAKA, "species_range") is None
    assert sites.not_applicable(sites.GREAT_NICOBAR, "accreditation") is None
    assert sites.not_applicable(sites.CUDDALORE, "omission") is None


# ---------- statewide species-list rule (and the mobile rule still applies) ----------

def statewide_gbif(region, asked=None, cls=BIRD):
    def gbif_get(path, **params):
        if asked is not None:
            asked.append(params.get("gadmGid"))
        if path == "/species/match":
            return cls
        if path == "/enumeration/country":
            return []
        if params.get("facet") == "country":
            return {"count": 0, "facets": []}
        if params.get("gadmGid") == "IND.16_1":
            return {"count": region}
        if params["taxonKey"] == 212:
            return {"count": 23000}
        return {"count": 0 if "geoDistance" in params else 5000}
    return gbif_get


def test_species_recorded_in_the_state_is_not_flagged(monkeypatch):
    asked = []
    mammal = {**BIRD, "class": "Mammalia"}                 # not mobile: the state rule still clears
    monkeypatch.setattr(species, "gbif_get", statewide_gbif(37, asked, mammal))
    r = species.check_species("Testus fakeus", 15.0, 76.25, archipelago_gadm=None,
                              region_gadm="IND.16_1", region_label="Karnataka")
    assert r["grade"] is None and r["region_count"] == 37
    assert "37 in Karnataka" in r["evidence"] and "Andaman" not in r["evidence"]
    assert species.ANDAMAN_NICOBAR_GADM not in asked


def test_species_absent_from_state_and_region_stays_red(monkeypatch):
    monkeypatch.setattr(species, "gbif_get", statewide_gbif(0))
    r = species.check_species("Testus fakeus", 15.0, 76.25, archipelago_gadm=None,
                              region_gadm="IND.16_1")
    assert r["grade"] == "RED"


def test_mobile_rule_still_applies_when_absent_from_state(monkeypatch):
    inner = statewide_gbif(0)

    def gbif_get(path, **params):
        if params.get("geoDistance", "").endswith(",3000km"):
            return {"count": 73}
        return inner(path, **params)
    monkeypatch.setattr(species, "gbif_get", gbif_get)
    r = species.check_species("Testus fakeus", 15.0, 76.25, archipelago_gadm=None,
                              region_gadm="IND.16_1")
    assert r["grade"] == "GREY" and r["mobile_region"]


def test_species_pages_pass_the_state_rule(monkeypatch):
    from tigerfish import species_pages as sp
    seen = []

    def check(name, lat, lon, archipelago_gadm="unset", region_gadm=None, region_label=None):
        seen.append((lat, lon, archipelago_gadm, region_gadm))
        return {"matched_name": name, "name_note": None, "grade": None, "local_count": 1,
                "regional_count": 1, "far_count": None, "global_count": 1, "class_name": "Aves",
                "evidence": "", "gbif_links": {}, "region_count": None}
    monkeypatch.setattr(sp.species, "check_species", check)
    monkeypatch.setattr(sp, "confirm", lambda name: MATCH)
    text = ("Table IV-20 List of birds recorded in Karnataka\nScientific Name\n"
            "Accipiter virgatus\nAviceda leuphotes\nIxobrychus flavicollis\nDicrurus macrocercus")
    out = sp.check_pages([{"page": 172, "text": text}], use_vision=False, profile=KA)
    assert seen and all(s == (15.0, 76.25, None, "IND.16_1") for s in seen)
    assert out["region_rule"] == "IND.16_1" and out["archipelago_rule"] is None
    assert out["rows"][0]["location"] == "Karnataka"


# ---------- pairing: the report's own places ----------

def test_pairing_uses_the_reports_places_and_restores_the_defaults():
    gn_prompt = pairing.KEY_PROMPT
    try:
        pairing.use_profile(KA)
        assert "Nicobar" not in pairing.KEY_PROMPT and "Galathea" not in pairing.KEY_PROMPT
        assert '"karnataka"' in pairing.KEY_PROMPT and "Asian Development Bank" in pairing.KEY_PROMPT
        key = {"measure": "road length", "place": "", "reading": False, "station": ""}
        claim = {"unit": "km", "value": 5178, "subject": "road length", "type": "quantity"}
        assert pairing.cell_of(claim, key)[2] == "project roads"
        assert pairing.place_relation("", "project roads") == "any"
        assert pairing.place_relation("karnataka", "project roads") == "nested"
        assert pairing.place_relation("karnataka", "ballari") is None
    finally:
        pairing.use_profile(None)
    assert pairing.KEY_PROMPT is gn_prompt                     # cached GN/Cuddalore keys still hit
    assert pairing.MERGED_PLACE == "great nicobar island"
    assert pairing.use_profile(sites.profile(sites.CUDDALORE)) is None and pairing.KEY_PROMPT is gn_prompt


# ---------- status tables: "SN" headers and WPAA / WLPAA columns ----------

KA_TABLE = "\n".join(["SN", "Common Name", "Scientific Name", "IUCN v.2025-2", "WPAA", "2022",
                      "28.", "Besra", "Accipiter virgatus", "Least Concern", "I",
                      "71.", "Bristled Grassbird", "Schoenicola striatus", "Vulnerable", "I",
                      "73.", "Bronzed Drongo", "Dicrurus aeneus", "Least concern", "0",
                      "78.", "Gundia Leaping Frog", "Indirana gundia", "Near", "Threatened", "NA"])


def test_status_parser_reads_the_karnataka_tables_and_never_guesses():
    t = status_pages.tables_from_pages([{"page": 172, "text": KA_TABLE}])
    assert len(t) == 1 and "IWPA Schedule" in t[0]["columns"]
    by = {r["Scientific name"]: r for r in t[0]["rows"]}
    assert by["Schoenicola striatus"]["IUCN"] == "Vulnerable" and by["Schoenicola striatus"]["IWPA Schedule"] == "I"
    assert by["Dicrurus aeneus"]["IWPA Schedule"] is None           # "0" is not a schedule
    assert by["Indirana gundia"]["IUCN"] is None                    # "Near" / "Threatened" split: not guessed


# ---------- the not-applicable scripts make no calls ----------

def refuse(*a, **k):
    raise AssertionError("no API call expected")


@pytest.mark.parametrize("script, out_name", [("check_omissions", "omissions.json"),
                                              ("check_accreditation", "accreditation.json")])
def test_not_applicable_scripts_write_a_reason_and_call_nothing(monkeypatch, tmp_path, script, out_name):
    mod = __import__(script)
    import check_omissions
    monkeypatch.setattr(check_omissions, "ROOT", tmp_path)
    monkeypatch.setattr(llm, "chat", refuse)
    monkeypatch.setattr(species, "gbif_get", refuse)
    monkeypatch.setattr(mod.pdf_reader, "read_pdf", refuse)
    monkeypatch.setattr(sys, "argv", [script, str(ROOT / "data" / "eias" / f"{sites.KARNATAKA}.pdf")])
    mod.main()
    saved = json.loads((tmp_path / "cache" / "findings" / sites.KARNATAKA / out_name).read_text(encoding="utf-8"))
    assert saved["grade"] == "NOT_APPLICABLE" and saved["not_applicable"]


def test_export_readers_turn_not_applicable_into_no_finding(monkeypatch, tmp_path):
    folder = tmp_path / "cache" / "findings" / sites.KARNATAKA
    folder.mkdir(parents=True)
    for name in ("omissions.json", "accreditation.json"):
        (folder / name).write_text(json.dumps({"not_applicable": "x", "grade": "NOT_APPLICABLE"}))
    monkeypatch.setattr(ex, "ROOT", tmp_path)
    pdf = Path(f"{sites.KARNATAKA}.pdf")
    assert ex.omission_findings(pdf) == [] and ex.accreditation_findings(pdf) == []
    notices = ex.not_applicable_notices(KA)
    assert len(notices) == 4 and all(sites.NOT_APPLICABLE in n for n in notices)


def test_karnataka_export_waits_for_the_audit(monkeypatch, tmp_path):
    monkeypatch.setattr(ex, "RESULTS", tmp_path)
    assert ex.karnataka() is None


# ---------- letter: Kannada, and addressed to KRRDA / ADB ----------

def fake_kannada(mangle_numbers=False):
    def chat(model, messages, **kw):
        assert model == llm.LIGHTNING and kw.get("thinking") is False
        assert "Kannada" in messages[0]["content"]
        texts = json.loads(messages[1]["content"])["texts"]
        out = ["ಕನ್ನಡ: " + t for t in texts]
        if mangle_numbers:
            out = [t.replace("392", "೩೯೨") for t in out]
        return json.dumps({"translations": out}, ensure_ascii=False)
    return chat


def test_kannada_offered_only_for_karnataka():
    opts = letter.language_options(KA_REPORT)
    assert opts["Kannada"] == "kn" and opts["English + Kannada"] == "both_kn"
    assert "Tamil" not in opts and letter.regional_languages(KA_REPORT) == ["Kannada"]
    assert not letter.offers_kannada({"title": "Cuddalore Port EIA", "state": "Tamil Nadu"})
    assert "Kannada" not in letter.language_options({"title": "Great Nicobar EIA"})


def test_kannada_letter_font_tag_and_machine_translated_label(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_kannada())
    body = letter.build_letter(KA_REPORT, FINDINGS, "both_kn", addressee="lender")
    doc = Document(io.BytesIO(body))
    runs = [r for p in doc.paragraphs for r in p.runs if letter.KANNADA.search(r.text)]
    assert runs
    for r in runs:
        rpr = r._element.rPr
        assert rpr.find(qn("w:rFonts")).get(qn("w:cs")) == letter.KANNADA_FONT == "Nirmala UI"
        assert rpr.find(qn("w:lang")).get(qn("w:val")) == "kn-IN"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Machine-translated; please review." in text
    assert text.index("Sir / Madam,") < text.index("ಕನ್ನಡ: Sir / Madam,")   # English first
    for q in FINDINGS[0]["quotes"]:
        assert f"“{q['text']}”" in text                                    # quotes never translated


def test_kannada_translation_that_changes_numbers_is_kept_in_english(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_kannada(mangle_numbers=True))
    text = doc_text(letter.build_letter(KA_REPORT, FINDINGS, "kn"))
    assert "The report states (page 392):" in text and "೩೯೨" not in text


def test_other_prompts_unchanged_so_cached_letters_still_hit():
    assert letter.TARGETS["hi"][0] is letter.TRANSLATE_PROMPT
    assert letter.TARGETS["ta"][0] is letter.TAMIL_PROMPT
    assert "Kannada" not in letter.TRANSLATE_PROMPT + letter.TAMIL_PROMPT


def test_lender_letter_goes_to_krrda_and_adb_not_the_eac():
    text = doc_text(letter.build_letter(KA_REPORT, FINDINGS, "en", addressee="lender"))
    assert "Karnataka Rural Roads Development Agency" in text and "Asian Development Bank" in text
    assert "Expert Appraisal Committee" not in text and "EIA Notification" not in text


def test_rti_pack_skips_the_public_hearing_letter():
    z = zipfile.ZipFile(io.BytesIO(rti.build_pack(KA_REPORT, FINDINGS)))
    assert "1_public_consultation_comments.docx" not in z.namelist()
    readme = z.read("README.txt").decode("utf-8")
    assert sites.NOT_APPLICABLE in readme and "2_RTI_applications.docx" in readme
    gn = zipfile.ZipFile(io.BytesIO(rti.build_pack({"title": "GN", "source": "x"}, FINDINGS)))
    assert "1_public_consultation_comments.docx" in gn.namelist()
