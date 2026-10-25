"""Letter languages with the translation model mocked: quotes, page numbers and links stay."""

import io
import json

import pytest
from docx import Document
from docx.oxml.ns import qn

from tigerfish import llm
from tigerfish.ui import letter

REPORT = {"title": "Great Nicobar EIA (2021 draft)", "source": "Draft EIA, December 2021"}
FINDINGS = [{
    "check_name": "Internal contradiction", "title": "Pages 392 and 978 disagree",
    "quotes": [{"page": 392, "text": "As per the ZSI report, no coral exists in the bay portion of the Galathea Bay."},
               {"page": 978, "text": "Coral reef exists in the proposed site for Transshipment / Port"}],
    "why": "Page 392 says no coral in the bay; page 978 says a coral reef exists at the port site.",
    "evidence": [{"label": "GBIF species page", "value": "Macaca fascicularis",
                  "url": "https://www.gbif.org/species/2436603"}],
    "clear": "A statement that the two passages cover different areas.",
}]


def fake_hindi(mangle_numbers=False):
    """LIGHTNING stand-in: 'translates' by prefixing Devanagari, keeping the text (and numbers)."""
    def chat(model, messages, **kw):
        assert model == llm.LIGHTNING and kw.get("thinking") is False
        texts = json.loads(messages[1]["content"])["texts"]
        out = []
        for t in texts:
            hi = "हिंदी: " + t
            if mangle_numbers:
                hi = hi.replace("392", "३९२")          # Devanagari digits: must be rejected
            out.append(hi)
        return json.dumps({"translations": out}, ensure_ascii=False)
    return chat


def doc_text(body):
    return "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)


def test_hindi_keeps_quotes_pages_and_links(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_hindi())
    text = doc_text(letter.build_letter(REPORT, FINDINGS, "hi"))
    assert text.splitlines()[0] == "Machine-translated; please review."
    for q in FINDINGS[0]["quotes"]:
        assert f"“{q['text']}”" in text                    # quotes untouched
        assert f"हिंदी: The report states (page {q['page']}):" in text
    assert "https://www.gbif.org/species/2436603" in text
    assert "हिंदी: Observation:" in text                              # our text translated


def test_quotes_are_never_sent_for_translation(monkeypatch):
    sent = []

    def chat(model, messages, **kw):
        sent.append(messages[1]["content"])
        return fake_hindi()(model, messages, **kw)

    monkeypatch.setattr(llm, "chat", chat)
    letter.build_letter(REPORT, FINDINGS, "hi")
    joined = "\n".join(sent)
    for q in FINDINGS[0]["quotes"]:
        assert q["text"] not in joined
    assert "Macaca fascicularis" not in joined                         # evidence value is data


def test_changed_numbers_fall_back_to_english(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_hindi(mangle_numbers=True))
    text = doc_text(letter.build_letter(REPORT, FINDINGS, "hi"))
    assert "The report states (page 392):" in text                     # English kept
    assert "३९२" not in text


def test_both_languages_in_one_document(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_hindi())
    text = doc_text(letter.build_letter(REPORT, FINDINGS, "both"))
    assert text.index("Sir / Madam,") < text.index("Machine-translated; please review.")
    assert text.count(FINDINGS[0]["quotes"][0]["text"]) == 2


def test_english_needs_no_model(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("no model call expected")
    monkeypatch.setattr(llm, "chat", boom)
    assert "Machine-translated" not in doc_text(letter.build_letter(REPORT, FINDINGS, "en"))


def test_devanagari_runs_use_a_complex_script_font(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_hindi())
    doc = Document(io.BytesIO(letter.build_letter(REPORT, FINDINGS, "hi")))
    runs = [r for p in doc.paragraphs for r in p.runs if "हिंदी" in r.text]
    assert runs
    fonts = runs[0]._element.rPr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:cs")) == letter.DEVANAGARI_FONT
    assert fonts.get(qn("w:eastAsia")) == letter.DEVANAGARI_FONT


def test_translation_unavailable(monkeypatch):
    def no_key(*a, **k):
        raise RuntimeError("NEBIUS_API_KEY is missing")
    monkeypatch.setattr(llm, "chat", no_key)
    with pytest.raises(letter.TranslationUnavailable):
        letter.build_letter(REPORT, FINDINGS, "hi")


def test_echoed_english_is_rejected(monkeypatch):
    """A model that copies the English back must not count as a translation."""
    def echo(model, messages, **kw):
        texts = json.loads(messages[1]["content"])["texts"]
        return json.dumps({"translations": texts})
    monkeypatch.setattr(llm, "chat", echo)
    assert letter.translate(["Sir / Madam,"]) == {}


def test_broken_chunk_is_split_and_retried(monkeypatch):
    calls = []

    def chat(model, messages, **kw):
        texts = json.loads(messages[1]["content"])["texts"]
        calls.append(len(texts))
        if len(texts) > 2:
            return "not json at all"
        return json.dumps({"translations": ["हिंदी: " + t for t in texts]}, ensure_ascii=False)

    monkeypatch.setattr(llm, "chat", chat)
    texts = [f"Point {i} on page {i + 10}" for i in range(4)]
    out = letter.translate(texts)
    assert len(out) == 4 and all(v.startswith("हिंदी") for v in out.values())
