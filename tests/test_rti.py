"""RTI + public hearing pack: zip contents, EIA Notification 2006 wording, one RTI question
per confirmed finding, neutral wording, placeholders, never the EP Act Section 5."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from tigerfish.ui import data, rti

ROOT = Path(__file__).resolve().parent.parent
BANNED = ("lie", "liar", "fraud", "fake", "fabricat", "dishonest", "cheat", "illegal", "violat")
pytestmark = pytest.mark.skipif(not (ROOT / "results" / "index.json").exists(), reason="no results/")


def docx_text(body):
    return "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)


@pytest.fixture
def gn():
    return data.load("great_nicobar_2021_draft")


@pytest.fixture
def pack(gn):
    picked = [f for f in gn["findings"] if f["id"] in ("pr-06", "sp-p954-001", "ac-01")
              or f.get("known_issue") == "gn-001"]
    with zipfile.ZipFile(io.BytesIO(rti.build_pack(gn, picked))) as z:
        return {n: z.read(n) for n in z.namelist()}, picked


def test_pack_has_readme_letter_and_rti(pack):
    files, _ = pack
    assert set(files) == {"README.txt", "1_public_consultation_comments.docx", "2_RTI_applications.docx"}
    assert "DRAFT, please review" in files["README.txt"].decode()


def test_letter_is_for_the_public_consultation_under_eia_notification_2006(pack):
    text = docx_text(pack[0]["1_public_consultation_comments.docx"])
    first = text.strip().splitlines()[0]
    assert first.startswith("DRAFT, please review")
    assert "EIA Notification, 2006" in text and "S.O. 1533(E)" in text
    assert "public consultation" in text and "[Name]" in text and "[address]" in text
    assert "page 392" in text and "page 978" in text


def test_one_rti_application_per_confirmed_finding(pack):
    files, picked = pack
    text = docx_text(files["2_RTI_applications.docx"])
    assert text.count("Section 6 of the Right to Information Act, 2005") == len(picked)
    assert text.count("DRAFT, please review") == len(picked)
    assert text.count("Application ") >= len(picked)
    assert "[Name]" in text and "[Address]" in text and "Public Information Officer" in text


def test_baseline_question_asks_for_the_survey_records(gn):
    f = next(f for f in gn["findings"] if f["id"] == "pr-06")
    q = rti.question(f)
    assert "survey records" in q and "page 292" in q and "3 months" in q
    assert "nine days" in q                               # the report's own description


def test_species_question_names_the_place_once(gn):
    f = next(f for f in gn["findings"] if f["id"] == "sp-p954-001")
    q = rti.question(f)
    assert "Panulirus interruptus" in q and "for Campbell Bay." in q and q.count("954") == 1


def test_questions_for_each_check_are_neutral_and_specific():
    for r in data.reports():
        report = data.load(r["id"])
        for f in report["findings"]:
            q = rti.question(f).lower()
            assert q.startswith(("certified copies", "copies", "the source")), (f["id"], q)
            words = set(q.replace(".", " ").replace(",", " ").split())
            assert not any(w.startswith(b) for w in words for b in BANNED), (f["id"], q)


def test_never_cites_environment_protection_act_section_5(gn):
    body = rti.build_pack(gn, gn["findings"])
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        texts = [z.read("README.txt").decode()] + [docx_text(z.read(n)) for n in z.namelist()
                                                    if n.endswith(".docx")]
    for t in texts:
        assert "Environment (Protection) Act" not in t and "Section 5" not in t


def test_committee_letter_is_unchanged_by_default(gn):
    text = docx_text(rti.letter.build_letter(gn, gn["findings"][:1]))
    assert "Expert Appraisal Committee" in text and "DRAFT" not in text
