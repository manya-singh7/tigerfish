"""Tamil letter (Tamil Nadu reports) on the same machine-translation path as Hindi; model mocked."""

import io
import json

from docx import Document
from docx.oxml.ns import qn

from tigerfish import llm
from tigerfish.ui import letter
from tests.test_letter import FINDINGS, REPORT, doc_text

TN_REPORT = {"title": "Cuddalore Port EIA (2026 draft)", "source": "Draft EIA", "state": "Tamil Nadu"}


def fake_tamil(mangle_numbers=False):
    def chat(model, messages, **kw):
        assert model == llm.LIGHTNING and kw.get("thinking") is False
        assert "Tamil" in messages[0]["content"]
        texts = json.loads(messages[1]["content"])["texts"]
        out = []
        for t in texts:
            ta = "தமிழ்: " + t
            out.append(ta.replace("392", "௩௯௨") if mangle_numbers else ta)
        return json.dumps({"translations": out}, ensure_ascii=False)
    return chat


def test_tamil_offered_only_for_tamil_nadu_reports():
    assert letter.offers_tamil(TN_REPORT)
    assert letter.offers_tamil({"title": "Port EIA, Cuddalore, Tamil Nadu"})
    assert not letter.offers_tamil(REPORT)
    assert "Tamil" in letter.language_options(TN_REPORT)
    assert letter.language_options(REPORT) == {"English": "en", "Hindi": "hi", "Both": "both"}


def test_tamil_keeps_quotes_pages_links_and_says_machine_translated(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_tamil())
    text = doc_text(letter.build_letter(TN_REPORT, FINDINGS, "ta"))
    assert text.splitlines()[0] == "Machine-translated; please review."
    for q in FINDINGS[0]["quotes"]:
        assert f"“{q['text']}”" in text                       # quotes never translated
        assert f"தமிழ்: The report states (page {q['page']}):" in text
    assert "https://www.gbif.org/species/2436603" in text


def test_tamil_uses_a_tamil_font_and_language_tag(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_tamil())
    doc = Document(io.BytesIO(letter.build_letter(TN_REPORT, FINDINGS, "both_ta")))
    runs = [r for p in doc.paragraphs for r in p.runs if letter.TAMIL.search(r.text)]
    assert runs
    for r in runs:
        rpr = r._element.rPr
        assert rpr.find(qn("w:rFonts")).get(qn("w:cs")) == letter.TAMIL_FONT
        assert rpr.find(qn("w:lang")).get(qn("w:val")) == "ta-IN"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.index("Sir / Madam,") < text.index("தமிழ்: Sir / Madam,")   # English first


def test_tamil_translation_that_changes_numbers_is_kept_in_english(monkeypatch):
    monkeypatch.setattr(llm, "chat", fake_tamil(mangle_numbers=True))
    text = doc_text(letter.build_letter(TN_REPORT, FINDINGS, "ta"))
    assert "The report states (page 392):" in text and "௩௯௨" not in text


def test_hindi_prompt_unchanged_so_cached_hindi_letters_still_hit():
    assert letter.TARGETS["hi"][0] is letter.TRANSLATE_PROMPT
    assert "Hindi" in letter.TRANSLATE_PROMPT and "Tamil" not in letter.TRANSLATE_PROMPT
