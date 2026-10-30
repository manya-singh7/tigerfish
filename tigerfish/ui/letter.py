"""Draft letter (.docx) from confirmed findings: neutral, factual, with page citations.

Languages: "en", "hi" (Hindi), "both" (English, then Hindi), "ta" (Tamil) or "both_ta"
(English, then Tamil), "kn" (Kannada) or "both_kn" (English, then Kannada). Tamil is offered
for Tamil Nadu reports (offers_tamil), Kannada for Karnataka reports (offers_kannada). All are
machine-translated by LIGHTNING (thinking off, cached) on the same path, and only OUR text is
translated: quotes from the report, evidence values, page numbers and links stay exactly as
they are. A translated segment that loses or changes a number or link is kept in English.
"""

import io
import json
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from tigerfish import authority
from tigerfish import claims as claims_mod
from tigerfish import llm

ACCENT = RGBColor(0x0F, 0x5E, 0x5E)
DEVANAGARI_FONT = "Nirmala UI"     # ships with Windows; Mangal is the older alternative
TAMIL_FONT = "Nirmala UI"          # covers Tamil too (Windows); Latha is the older alternative
KANNADA_FONT = "Nirmala UI"        # covers Kannada too (all 90 glyphs of U+0C80-0CFF); Tunga is older
LATIN_FONT = "Calibri"
MACHINE_NOTE = "Machine-translated; please review."
TRANSLATE_CHUNK = 8                # segments per LIGHTNING call
URL = re.compile(r"https?://\S+")
DEVANAGARI = re.compile(r"[ऀ-ॿ]")
TAMIL = re.compile(r"[஀-௿]")
KANNADA = re.compile(r"[ಀ-೿]")
NUM = re.compile(r"\d[\d,]*(?:\.\d+)?")

TRANSLATE_PROMPT = """You translate English into formal, plain Hindi (Devanagari script) for \
a letter to a government committee.
Input: JSON {"texts": [...]}. Output: JSON only, {"translations": [...]}, with the Hindi \
translation of each text, same order and same count. Write the translation itself, in \
Devanagari; do not copy the English and do not add numbering.
Keep every number, page number, URL, species name and report title exactly as written \
(use Western digits 0-9) and keep placeholders in square brackets like [Name] unchanged.
Example: {"texts": ["Sir / Madam,", "Clarification requested: page 12"]} -> \
{"translations": ["महोदय / महोदया,", "स्पष्टीकरण का अनुरोध: पृष्ठ 12"]}"""


TAMIL_PROMPT = """You translate English into formal, plain Tamil (Tamil script) for a letter to a government committee.
Input: JSON {"texts": [...]}. Output: JSON only, {"translations": [...]}, with the Tamil translation of each text, same order and same count. Write the translation itself, in Tamil script; do not copy the English and do not add numbering.
Keep every number, page number, URL, species name and report title exactly as written (use Western digits 0-9) and keep placeholders in square brackets like [Name] unchanged.
Example: {"texts": ["Sir / Madam,", "Clarification requested: page 12"]} -> {"translations": ["ஐயா / அம்மையீர்,", "விளக்கம் கோரப்படுகிறது: பக்கம் 12"]}"""

KANNADA_PROMPT = """You translate English into formal, plain Kannada (Kannada script) for a letter to a government committee.
Input: JSON {"texts": [...]}. Output: JSON only, {"translations": [...]}, with the Kannada translation of each text, same order and same count. Write the translation itself, in Kannada script; do not copy the English and do not add numbering.
Keep every number, page number, URL, species name and report title exactly as written (use Western digits 0-9) and keep placeholders in square brackets like [Name] unchanged.
Example: {"texts": ["Sir / Madam,", "Clarification requested: page 12"]} -> {"translations": ["ಮಾನ್ಯರೇ,", "ಸ್ಪಷ್ಟೀಕರಣ ಕೋರಲಾಗಿದೆ: ಪುಟ 12"]}"""

# language code -> (prompt, script regex, font, Word language tag)
TARGETS = {"hi": (TRANSLATE_PROMPT, DEVANAGARI, DEVANAGARI_FONT, "hi-IN"),
           "ta": (TAMIL_PROMPT, TAMIL, TAMIL_FONT, "ta-IN"),
           "kn": (KANNADA_PROMPT, KANNADA, KANNADA_FONT, "kn-IN")}
TAMIL_STATES = {"tamil nadu", "puducherry"}
KANNADA_STATES = {"karnataka"}


def offers_tamil(report):
    """Tamil letters are offered for reports from Tamil Nadu: the report's "state", or its
    title / source naming the state."""
    if str(report.get("state") or "").strip().lower() in TAMIL_STATES:
        return True
    text = f"{report.get('title', '')} {report.get('source', '')}".lower()
    return any(state in text for state in TAMIL_STATES)


def offers_kannada(report):
    """Kannada letters are offered for reports from Karnataka (the report's "state")."""
    return str(report.get("state") or "").strip().lower() in KANNADA_STATES


def regional_languages(report):
    """Names of the regional languages offered for this report (for the picker's help text)."""
    return [name for name, offered in (("Tamil", offers_tamil(report)),
                                       ("Kannada", offers_kannada(report))) if offered]


def language_options(report):
    """{label: code} for the letter-language picker."""
    out = {"English": "en", "Hindi": "hi", "Both": "both"}
    if offers_tamil(report):
        out.update({"Tamil": "ta", "English + Tamil": "both_ta"})
    if offers_kannada(report):
        out.update({"Kannada": "kn", "English + Kannada": "both_kn"})
    return out


class TranslationUnavailable(RuntimeError):
    pass


# ---------- letter content (English blocks) ----------

def where(q):
    return f"page {q['page']}" if q.get("page") else q.get("where", "the report")


DRAFT_NOTE = ("DRAFT, please review. Check every point against the report, fill in the "
              "[placeholders] and remove this line before sending.")
ADDRESSEES = ("committee", "public_consultation", "lender")


def blocks(report, findings, addressee="committee", draft=False):
    """The letter as blocks. kind: text (ours, translatable), quote / data (never translated),
    heading / subject (ours, translatable, styled), break.
    addressee: "committee" (the Expert Appraisal Committee) or "public_consultation" (written
    comments in the public consultation under the EIA Notification, 2006) or "lender" (a
    report prepared for a development bank, e.g. ADB: to the borrower's agency, copied to the
    bank; report["letter_to"] names them)."""
    if addressee not in ADDRESSEES:
        raise ValueError(f"addressee must be one of {ADDRESSEES}")
    out = [("subject", DRAFT_NOTE)] if draft else []
    out.append(("data", date.today().strftime("%d %B %Y")))
    if addressee == "lender":
        out += [*authority.lender_letter_blocks(report, "To,\n" + report.get(
                    "letter_to", "[the agency that prepared the report]\n[the financing institution]")
                    + "\n[address]"),
                ("subject", f"Subject: Observations on {report['title']}, request for clarification"),
                ("text", "Sir / Madam,"),
                ("text", f"I have reviewed {report['title']} ({report.get('source', '')}). The points "
                         f"below are statements in the report that appear inconsistent with each other "
                         f"or with public records. Each point cites the page and the exact text. I "
                         f"request that the agency that prepared the report and the financing "
                         f"institution verify these points and clarify or correct them before the "
                         f"report is finalised and relied upon.")]
    elif addressee == "public_consultation":
        out += [*authority.to_blocks(report, "hearing", "To,\nThe Member Secretary\n[State Pollution "
                                     "Control Board / Union Territory Pollution Control Committee "
                                     "conducting the public hearing]\n[address]"),
                ("subject", f"Subject: Written comments in the public consultation on "
                            f"{report['title']}, under the EIA Notification, 2006"),
                ("text", "Sir / Madam,"),
                ("text", f"I submit these written comments in the public consultation (Stage 3) "
                         f"under the Environment Impact Assessment Notification, 2006 (S.O. 1533(E), "
                         f"14 September 2006), on {report['title']} ({report.get('source', '')}). "
                         f"The points below are statements in the draft report that appear "
                         f"inconsistent with each other or with public records. Each point cites "
                         f"the page and the exact text. I request that these comments be placed on "
                         f"record with the responses received in the public consultation, and that "
                         f"the project proponent and the consultant verify these points and clarify "
                         f"or correct them in the final report.")]
    else:
        out += [*authority.to_blocks(report, "clearance", "To,\nThe Member Secretary\nExpert Appraisal "
                                     "Committee\nMinistry of Environment, Forest and Climate Change\n[address]"),
                ("subject", f"Subject: Observations on {report['title']}, request for clarification"),
                ("text", "Sir / Madam,"),
                ("text", f"I have reviewed {report['title']} ({report.get('source', '')}). The points "
                         f"below are statements in the report that appear inconsistent with each other "
                         f"or with public records. Each point cites the page and the exact text. I "
                         f"request that the project proponent and the consultant verify these points "
                         f"and clarify or correct them before the report is relied upon.")]
    for i, f in enumerate(findings, 1):
        out.append(("heading", f"{i}. {f['check_name']}: {f['title']}"))
        for q in f["quotes"]:
            out.append(("quote_label", f"The report states ({where(q)}):"))
            out.append(("quote", f"“{q['text']}”"))
        out.append(("text", f"Observation: {f['why']}"))
        for e in f["evidence"]:
            out.append(("evidence", e["label"], str(e.get("value") or ""), e.get("url")))
        out.append(("text", f"Clarification requested: {f['clear']}"))
    out += [("text", "These observations were prepared with the help of TigerFish, an automated "
                     "accuracy checker, and each point was reviewed by me before inclusion. They "
                     "are offered so that the record can be verified and, where needed, corrected."),
            ("text", "Yours faithfully,\n\n[Name]\n[Address / organisation]\n[Contact]")]
    return out


def translatable(blks):
    """The English strings we translate (our own text only), in block order."""
    out = []
    for b in blks:
        if b[0] in ("text", "subject", "heading", "quote_label"):
            out.append(b[1])
        elif b[0] == "evidence":
            out.append(b[1])                    # the label only; the value is data
    return out


# ---------- translation ----------

def keeps_numbers_and_links(src, dst):
    """Every number and URL in the English text must appear unchanged in the translation."""
    return (set(NUM.findall(src)) <= set(NUM.findall(dst))
            and set(URL.findall(src)) <= set(URL.findall(dst)))


def is_translation(src, dst, lang="hi"):
    """Text in the target script (Devanagari / Tamil / Kannada) that differs from the English, keeping
    its numbers and links."""
    return (bool(TARGETS[lang][1].search(dst)) and dst.strip() != src.strip()
            and keeps_numbers_and_links(src, dst))


def translate(texts, lang="hi"):
    """{english: translation} via LIGHTNING (cached). Unsafe or missing translations are left
    out (the English stays). Raises TranslationUnavailable if the model can't be reached."""
    out = {}
    unique = list(dict.fromkeys(t for t in texts if t.strip()))
    for start in range(0, len(unique), TRANSLATE_CHUNK):
        _translate_chunk(unique[start:start + TRANSLATE_CHUNK], out, lang)
    return out


def _translate_chunk(chunk, out, lang="hi"):
    """Translate one chunk into `out`; if the reply is unusable, split it and try the halves
    (down to single texts), so one bad reply doesn't leave a whole chunk in English."""
    try:
        raw = claims_mod.chat_json(llm.LIGHTNING, [
            {"role": "system", "content": TARGETS[lang][0]},
            {"role": "user", "content": json.dumps({"texts": chunk}, ensure_ascii=False)}],
            max_tokens=4000)
    except Exception as e:                      # no key, network, ...
        raise TranslationUnavailable(str(e)) from e
    got = (raw or {}).get("translations") if isinstance(raw, dict) else None
    if not isinstance(got, list) or len(got) != len(chunk):
        if len(chunk) > 1:
            half = len(chunk) // 2
            _translate_chunk(chunk[:half], out, lang)
            _translate_chunk(chunk[half:], out, lang)
        return
    for src, dst in zip(chunk, got):
        if isinstance(dst, str) and is_translation(src, dst, lang):
            out[src] = dst.strip()


# ---------- rendering ----------

def set_font(run, name):
    """Set a run's font for Latin, East Asian and complex scripts (Devanagari is 'cs')."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def add_run(par, text, bold=False, italic=False, color=None):
    """A run with the right font: Devanagari / Tamil / Kannada text gets an Indic font and a
    language tag."""
    run = par.add_run(text)
    target = next((t for t in TARGETS.values() if t[1].search(text)), None)
    if target:
        set_font(run, target[2])
        lang = OxmlElement("w:lang")
        lang.set(qn("w:bidi"), target[3])
        lang.set(qn("w:val"), target[3])
        run._element.get_or_add_rPr().append(lang)
    else:
        set_font(run, LATIN_FONT)
    run.bold, run.italic = bold or None, italic or None
    if color:
        run.font.color.rgb = color
    return run


def render(doc, blks, tr=None):
    tr = tr or {}
    t = lambda s: tr.get(s, s)                      # noqa: E731
    for b in blks:
        kind = b[0]
        if kind == "data":
            add_run(doc.add_paragraph(), b[1])
        elif kind == "subject":
            add_run(doc.add_paragraph(), t(b[1]), bold=True, color=ACCENT)
        elif kind == "heading":
            add_run(doc.add_paragraph(), t(b[1]), bold=True)
        elif kind == "quote_label":
            par = doc.add_paragraph(style="List Bullet")
            add_run(par, t(b[1]) + " ", italic=True)
        elif kind == "quote":                          # appended to the label's paragraph
            add_run(doc.paragraphs[-1], b[1])          # never translated
        elif kind == "evidence":
            _, label, value, url = b
            par = doc.add_paragraph(style="List Bullet 2")
            add_run(par, f"{t(label)}: ")
            add_run(par, value)                        # data: kept as is
            if url:
                add_run(par, f" ({url})")
        else:
            add_run(doc.add_paragraph(), t(b[1]))


def preview_html(report, findings, addressee="committee"):
    """The English letter as HTML, for reading on the page before downloading it."""
    from tigerfish.ui.style import esc
    out = []
    for b in blocks(report, findings, addressee):
        kind, text = b[0], b[1]
        if kind == "evidence":
            out.append(f'<p class="tf-letter-ev">{esc(text)}: {esc(b[2])}</p>')
        elif kind == "quote":
            out.append(f'<blockquote class="tf-letter-quote">{esc(text)}</blockquote>')
        else:
            tag = {"subject": "b", "heading": "b", "quote_label": "i"}.get(kind)
            body = esc(text).replace("\n", "<br>")
            out.append(f"<p><{tag}>{body}</{tag}></p>" if tag else f"<p>{body}</p>")
    return f'<div class="tf-letter">{"".join(out)}</div>'


def build_letter(report, findings, language="en", addressee="committee", draft=False):
    """Return .docx bytes. language: "en", "hi", "both" (en + hi), "ta" or "both_ta" (en + ta),
    "kn" or "both_kn" (en + kn); addressee and draft: see blocks()."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    blks = blocks(report, findings, addressee, draft)
    target = {"hi": "hi", "both": "hi", "ta": "ta", "both_ta": "ta",
              "kn": "kn", "both_kn": "kn"}.get(language)
    if language in ("en", "both", "both_ta", "both_kn"):
        render(doc, blks)
    if target:
        tr = translate(translatable(blks) + [MACHINE_NOTE], target)
        if language in ("both", "both_ta", "both_kn"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        add_run(doc.add_paragraph(), MACHINE_NOTE, bold=True)
        if MACHINE_NOTE in tr:
            add_run(doc.add_paragraph(), tr[MACHINE_NOTE], italic=True)
        render(doc, blks, tr)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
