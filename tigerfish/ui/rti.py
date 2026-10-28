"""RTI + public hearing pack: a .zip with
  1. written comments for the EIA public consultation under the EIA Notification, 2006
     (the existing letter, addressed to the body conducting the public hearing), and
  2. draft RTI applications under Section 6 of the Right to Information Act, 2005,
     one application with one question per confirmed finding.
Both are drafts with [placeholders]. Wording is neutral: each question asks for the records
behind a statement, never says a statement is wrong. No model calls.
"""

import io
import re
import zipfile
from datetime import date

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from tigerfish import authority
from tigerfish.ui import letter

DRAFT_NOTE = ("DRAFT, please review. Check each question against the report, fill in the "
              "[placeholders] and remove this line before sending.")
README = """TigerFish: RTI and public hearing pack (DRAFT, please review)

1_public_consultation_comments.docx
   Written comments for the public consultation on the draft EIA report, under the
   Environment Impact Assessment Notification, 2006 (Stage 3). Send them to the body
   conducting the public hearing, before or at the hearing.

2_RTI_applications.docx
   Draft applications under Section 6 of the Right to Information Act, 2005, one per
   confirmed finding, each asking for the records behind one statement in the report.
   Address each to the Public Information Officer of the public authority that holds
   those records, and pay the application fee set by the applicable RTI rules.

Everything was drafted from findings a reviewer marked Confirm in TigerFish. Fill in every
[placeholder], check each point against the report, and remove the DRAFT lines before sending.
"""


def where(q):
    return f"page {q['page']}" if q.get("page") else (q.get("where") or "the report")


def _ev(f, start):
    e = next((e for e in f["evidence"] if e["label"].startswith(start)), None)
    return (e or {}).get("value")


def _short(text, words=14):
    w = str(text).split()
    return " ".join(w[:words]) + (" ..." if len(w) > words else "")


def _quotes(f):
    return "; ".join(f"{where(q)}: “{_short(q['text'], 25)}”" for q in f["quotes"])


def question(f):
    """One neutral request for records, for one finding."""
    check, qs = f.get("check"), f["quotes"]
    page_q = [q for q in qs if q.get("page")]
    first = (page_q or qs)[0]
    if check == "promise":
        topic = f["title"].split(":", 1)[0].strip().lower()
        measure = f["title"].split(":", 1)[1].replace("in the Terms of Reference", "").strip() \
            if ":" in f["title"] else ""
        described = _ev(f, "Report describes") or ""
        tor = next((q.get("where") for q in qs if not q.get("page")), "the Terms of Reference")
        return (f"Certified copies of the complete field survey records behind the {topic} "
                f"baseline described on {where(first)} of the report"
                + (f" (“{described}”)" if described else "")
                + ", including the dates and locations of each survey, the field data sheets, and "
                f"the names of the teams or agencies that carried them out; and copies of any "
                f"records of {topic} baseline data collected over the {measure} stated in the "
                f"{tor}.")
    if check == "contradiction":
        return (f"Copies of the survey records, data sheets or source documents on which the "
                f"following statements in the report are based: {_quotes(f)}; and any record "
                f"of how these statements relate to each other.")
    if check == "recycled_figure":
        return (f"Copies of the species lists or survey records on which each of the following "
                f"figures in the report is based, for each place it is given for: {_quotes(f)}.")
    if check == "species_range":
        place = f["title"].split(" at ", 1)[1] if " at " in f["title"] else "the project area"
        place = re.sub(r"\s*\(p\.\d+\)$", "", place)          # page is already cited
        return (f"Copies of the survey, sighting or specimen records (date, location, method and "
                f"observer) for {first['text']}, as listed on {where(first)} of the report for "
                f"{place}.")
    if check == "omission":
        species = f["title"].split(" (", 1)[0]
        return (f"Copies of any records showing whether {species} was surveyed for or considered "
                f"in the baseline studies for the project area, including the species lists "
                f"compiled during the field surveys.")
    if check == "accreditation":
        consultant = f["title"].split(":", 1)[0]
        return (f"Copies of the QCI/NABET accreditation certificate(s) of {consultant} that were "
                f"valid on the date of the report, showing the accredited sectors and categories, "
                f"as held on the file for this project.")
    if check in ("borrowed_text", "copy_paste"):
        return (f"The source documents used for the text on {where(first)} of the report that "
                f"begins “{_short(first['text'])}”, as held in the records for this project.")
    if check in ("iucn_status", "legal_schedule"):
        what = "IUCN Red List category" if check == "iucn_status" else "Wild Life (Protection) Act schedule"
        species = f["title"].split(":", 1)[0]
        return (f"The source, and its date, of the {what} given for {species} in {where(first)} "
                f"of the report.")
    return (f"Copies of the records on which the statement on {where(first)} of the report "
            f"(“{_short(first['text'], 25)}”) is based.")


def application_blocks(report, f):
    """One RTI application as (kind, text) blocks."""
    return [
        ("draft", DRAFT_NOTE),
        ("text", authority.rti_to(report, "To,\nThe Public Information Officer\n[Name of the public "
                 "authority that holds these records, for example the Ministry of Environment, Forest "
                 "and Climate Change (Impact Assessment Division), or the project proponent if it is "
                 "a public authority]\n[address]")),
        ("subject", "Subject: Application for information under Section 6 of the Right to "
                    "Information Act, 2005"),
        ("text", "Sir / Madam,"),
        ("text", "1. Name of the applicant: [Name]"),
        ("text", "2. Address for correspondence: [Address]\n   Phone / email (optional): [Contact]"),
        ("text", f"3. Particulars of the information requested, about {report['title']} "
                 f"({report.get('source', '')}):"),
        ("question", question(f)),
        ("text", "4. Period to which the information relates: [period, for example the period of "
                 "the baseline studies]"),
        ("text", "5. Fee: [the application fee under the applicable RTI rules, and how it is paid: "
                 "postal order / demand draft / online, with reference number]. [If applicable: "
                 "I belong to the Below Poverty Line category; a copy of the certificate is "
                 "attached.]"),
        ("text", "6. I would like the information as [certified photocopies / electronic copies "
                 "by email]."),
        ("text", "I am a citizen of India. If the information is held by another public "
                 "authority, I request that this application be transferred to it under "
                 "Section 6(3) of the Act."),
        ("text", f"Place: [place]\nDate: {date.today().strftime('%d %B %Y')}\n\n"
                 f"[Signature]\n[Name]"),
    ]


def build_rti(report, findings):
    """.docx bytes: one draft application per finding, each on its own page."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for i, f in enumerate(findings):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph().add_run(f"Application {i + 1} of {len(findings)}: {f['title']}").italic = True
        for kind, text in application_blocks(report, f):
            run = doc.add_paragraph().add_run(text)
            run.bold = kind in ("draft", "subject")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pack(report, findings):
    """.zip bytes: README, public consultation comments (.docx), RTI applications (.docx)."""
    buf = io.BytesIO()
    hearing_na = (report.get("not_applicable") or {}).get("public_consultation")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        if hearing_na:   # e.g. no EIA Notification public hearing: only the RTI drafts
            z.writestr("README.txt", README.replace(
                README[README.index("1_public"):README.index("2_RTI")],
                f"1_public_consultation_comments.docx: not applicable to this report.\n"
                f"   {hearing_na}\n\n"))
        else:
            z.writestr("README.txt", README)
            z.writestr("1_public_consultation_comments.docx",
                       letter.build_letter(report, findings, "en", addressee="public_consultation",
                                           draft=True))
        z.writestr("2_RTI_applications.docx", build_rti(report, findings))
    return buf.getvalue()
